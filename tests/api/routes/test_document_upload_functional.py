"""Multipart upload checks with real parsing, pipeline and local storages."""

import importlib
from io import BytesIO
import sys
from uuid import uuid4

import numpy as np
import pytest
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient

_original_argv = sys.argv[:]
try:
    sys.argv = sys.argv[:1]
    routes = importlib.import_module("lightrag.api.routers.document_routes")
finally:
    sys.argv = _original_argv

from lightrag import LightRAG  # noqa: E402
from lightrag.kg.folder_storage import FolderManager  # noqa: E402
from lightrag.kg.shared_storage import get_namespace_data  # noqa: E402
from lightrag.utils import EmbeddingFunc, Tokenizer  # noqa: E402

pytestmark = pytest.mark.offline


class CharacterTokenizer:
    def encode(self, text):
        return list(map(ord, text))

    def decode(self, tokens):
        return "".join(map(chr, tokens))


async def fake_embedding(texts, **kwargs):
    return np.ones((len(texts), 8))


async def fake_llm(*args, **kwargs):
    return "<|COMPLETE|>"


@pytest.fixture
async def upload_app(tmp_path, monkeypatch):
    input_dir = tmp_path / "inputs"
    monkeypatch.setenv("INPUT_DIR", str(input_dir))
    monkeypatch.setenv("LIGHTRAG_PARSER", "*:legacy")
    monkeypatch.setenv("LIGHTRAG_OCR_ENGINE", "auto")
    rag = LightRAG(
        working_dir=str(tmp_path / "storage"),
        workspace=f"upload-{uuid4().hex}",
        tokenizer=Tokenizer("test", CharacterTokenizer()),
        llm_model_func=fake_llm,
        embedding_func=EmbeddingFunc(embedding_dim=8, func=fake_embedding),
        vlm_process_enable=False,
    )
    await rag.initialize_storages()
    folders_kv = rag.key_string_value_json_storage_cls(
        namespace="doc_folders",
        workspace=rag.workspace,
        embedding_func=None,
    )
    await folders_kv.initialize()
    folders = FolderManager(folders_kv, rag.workspace)
    # Match create_app: folder filtering attaches a live manager to the RAG.
    rag.folder_manager = folders
    folder = await folders.create_folder("Upload checks")
    app = FastAPI()
    app.include_router(
        routes.create_document_routes(
            rag,
            routes.DocumentManager(str(input_dir)),
            api_key="upload-test-key",
            folder_manager=folders,
        )
    )
    try:
        async with AsyncClient(
            transport=ASGITransport(app=app),
            base_url="http://test",
            headers={"X-API-Key": "upload-test-key"},
        ) as client:
            yield client, rag, folder, input_dir
    finally:
        await rag.finalize_storages()
        await folders_kv.finalize()


@pytest.mark.parametrize(
    "filename",
    [
        "upload.txt",
        "上传测试.md",
        "data.csv",
        "data.json",
        "文档.docx",
        "原生.[native-F].docx",
        "report.pdf",
    ],
)
@pytest.mark.parametrize("use_folder", [False, True])
async def test_new_file_upload_is_processed_and_visible(
    upload_app, filename, use_folder
):
    client, rag, folder, input_dir = upload_app
    if filename.endswith(".docx"):
        from docx import Document

        document = Document()
        document.add_paragraph("Alice works at Example Labs. 这是一份文件上传测试。")
        stream = BytesIO()
        document.save(stream)
        content = stream.getvalue()
    elif filename.endswith(".pdf"):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {
                        NameObject("/F1"): DictionaryObject(
                            {
                                NameObject("/Type"): NameObject("/Font"),
                                NameObject("/Subtype"): NameObject("/Type1"),
                                NameObject("/BaseFont"): NameObject("/Helvetica"),
                            }
                        )
                    }
                )
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 20 200 Td (Alice works at Example Labs.) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
        output = BytesIO()
        writer.write(output)
        content = output.getvalue()
    else:
        content = b"Alice works at Example Labs. This is an upload test."
    response = await client.post(
        "/documents/upload",
        files={"file": (filename, content)},
        data={"folder_id": folder.id} if use_folder else {},
    )
    assert response.status_code == 200, response.text
    track_id = response.json()["track_id"]
    docs = await rag.doc_status.get_docs_by_track_id(track_id)
    assert len(docs) == 1
    doc_id, doc = next(iter(docs.items()))
    assert doc.status.value == "processed", doc.error_msg
    assert doc.chunks_list
    assert await rag.full_docs.get_by_id(doc_id)
    assert all(await rag.text_chunks.get_by_ids(doc.chunks_list))
    if use_folder:
        assert folder.id in doc.metadata["folder_ids"]
    listed = await client.post(
        "/documents/paginated",
        json={
            "folder_id": folder.id if use_folder else None,
            "status_filters": ["processed"],
        },
    )
    assert listed.status_code == 200, listed.text
    assert [item["id"] for item in listed.json()["documents"]] == [doc_id]
    assert list((input_dir / "__parsed__").rglob("*"))
    status = await get_namespace_data("pipeline_status", workspace=rag.workspace)
    assert status["pending_enqueues"] == 0
    assert not status["busy"]


async def test_missing_file_field_returns_validation_error(upload_app):
    client, _, _, _ = upload_app
    response = await client.post(
        "/documents/upload", files={"wrong_field": ("test.txt", b"test")}
    )
    assert response.status_code == 422
    assert response.json()["detail"][0]["loc"] == ["body", "file"]


@pytest.mark.parametrize("mode", ["normal", "fast", "parse_failed", "kg_failed"])
@pytest.mark.parametrize("callback_url", ["http://caller.example/finished", "http://user:password@caller.example/finished"])
async def test_upload_completion_callback(upload_app, monkeypatch, mode, callback_url):
    from lightrag.api import upload_callbacks

    client, rag, _, _ = upload_app
    received = []

    async def capture(url, payload, **kwargs):
        received.append((url, payload))

    async def fail_kg(*args, **kwargs):
        raise RuntimeError("KG test failure")

    monkeypatch.setattr(upload_callbacks, "send_upload_callback", capture)
    if mode == "kg_failed":
        monkeypatch.setattr(rag, "_process_extract_entities", fail_kg)
    response = await client.post(
        "/documents/upload",
        files={"file": ("callback.txt", b"" if mode == "parse_failed" else b"Callback document.")},
        data={"callback_url": callback_url, "fast_index": str(mode == "fast").lower()},
    )
    assert response.status_code == 200, response.text
    assert len(received) == 1
    url, payload = received[0]
    assert url == callback_url
    assert payload["track_id"] == response.json()["track_id"]
    assert payload["event_id"] == payload["track_id"]
    assert payload["filename"] == "callback.txt"
    assert payload["status"] == ("failed" if "failed" in mode else "processed")
    assert len(payload["documents"]) == 1
    doc = payload["documents"][0]
    if mode == "normal":
        assert doc["kg_status"] == "completed"
    elif mode == "fast":
        assert doc["kg_status"] == "skipped"
    else:
        assert doc["error"]


@pytest.mark.parametrize("url", ["invalid", "ftp://caller.example/file"])
async def test_invalid_callback_url_rejected_before_upload(upload_app, url):
    client, rag, _, input_dir = upload_app
    response = await client.post(
        "/documents/upload", files={"file": ("invalid-callback.txt", b"content")},
        data={"callback_url": url},
    )
    assert response.status_code == 422
    assert not list(input_dir.rglob("invalid-callback*"))
    status = await get_namespace_data("pipeline_status", workspace=rag.workspace)
    assert status.get("pending_enqueues", 0) == 0


async def test_callback_delivery_logs_visible_in_webui_status(upload_app, monkeypatch):
    import httpx
    from lightrag.api import upload_callbacks

    client, _, _, _ = upload_app
    requests = []

    def handler(request):
        requests.append(request)
        return httpx.Response(500 if len(requests) == 1 else 204)

    original = httpx.AsyncClient
    monkeypatch.setattr(
        upload_callbacks.httpx, "AsyncClient",
        lambda **kwargs: original(transport=httpx.MockTransport(handler), **kwargs),
    )
    response = await client.post(
        "/documents/upload", files={"file": ("callback-log.txt", b"Callback log document.")},
        data={"callback_url": "http://user:secret@caller.example/done?token=private-token"},
    )
    assert response.status_code == 200, response.text
    status_response = await client.get("/documents/pipeline_status")
    assert status_response.status_code == 200, status_response.text
    status = status_response.json()
    messages = [msg for msg in status["history_messages"] if msg.startswith("[callback_url]")]
    assert len(messages) == 5
    assert "Sending" in messages[0]
    assert "HTTP=500" in messages[1]
    assert "Retry scheduled" in messages[2]
    assert "Success" in messages[-1] and "HTTP=204" in messages[-1]
    assert all(response.json()["track_id"] in msg for msg in messages)
    assert all("secret" not in msg and "private-token" not in msg for msg in messages)
    assert len(status["history_message_timings"]) == len(status["history_messages"])
    assert status["history_message_timings"][-1]["time"]


@pytest.mark.parametrize(
    "filename, content",
    [
        ("empty.txt", b""),
        ("broken.pdf", b"invalid pdf"),
        ("broken.docx", b"invalid docx"),
    ],
)
async def test_failed_upload_stays_visible_in_selected_folder(
    upload_app, filename, content
):
    client, rag, folder, _ = upload_app
    response = await client.post(
        "/documents/upload",
        files={"file": (filename, content)},
        data={"folder_id": folder.id},
    )
    assert response.status_code == 200
    docs = await rag.doc_status.get_docs_by_track_id(response.json()["track_id"])
    assert len(docs) == 1
    doc_id, doc = next(iter(docs.items()))
    assert doc.status.value == "failed"
    assert doc.error_msg
    listed = await client.post(
        "/documents/paginated",
        json={"folder_id": folder.id, "status_filters": ["failed"]},
    )
    assert listed.status_code == 200
    assert [item["id"] for item in listed.json()["documents"]] == [doc_id]


async def test_invalid_native_heading_does_not_break_later_uploads(upload_app):
    """A parser validation failure must leave shared storage and workers usable."""
    from docx import Document

    client, rag, folder, _ = upload_app
    document = Document()
    document.add_heading("H" * 201, level=1)
    document.add_paragraph("This document has an invalid heading.")
    stream = BytesIO()
    document.save(stream)
    response = await client.post(
        "/documents/upload",
        files={"file": ("invalid-heading.[native-F].docx", stream.getvalue())},
        data={"folder_id": folder.id},
    )
    assert response.status_code == 200, response.text
    docs = await rag.doc_status.get_docs_by_track_id(response.json()["track_id"])
    assert len(docs) == 1
    failed_doc = next(iter(docs.values()))
    assert failed_doc.status.value == "failed"
    assert "Heading too long" in failed_doc.error_msg
    assert folder.id in failed_doc.metadata["folder_ids"]

    status = await get_namespace_data("pipeline_status", workspace=rag.workspace)
    assert not status["busy"]
    assert status["pending_enqueues"] == 0

    response = await client.post(
        "/documents/upload",
        files={"file": ("after-invalid.txt", b"Normal upload after DOCX validation failure.")},
    )
    assert response.status_code == 200, response.text
    docs = await rag.doc_status.get_docs_by_track_id(response.json()["track_id"])
    assert len(docs) == 1
    assert next(iter(docs.values())).status.value == "processed"


@pytest.mark.parametrize("suffix,format", [("png", "PNG"), ("jpg", "JPEG"), ("pdf", "PDF")])
async def test_vlm_ocr_upload_reaches_processed(upload_app, monkeypatch, suffix, format):
    from PIL import Image, ImageDraw
    from unittest.mock import AsyncMock
    import shutil

    if suffix == "pdf" and not shutil.which("pdftoppm"):
        pytest.skip("Poppler is required to exercise real scanned-PDF rendering")
    client, rag, folder, _ = upload_app
    monkeypatch.setenv("LIGHTRAG_OCR_ENGINE", "vlm")
    rag.vlm_process_enable = True
    vlm = AsyncMock(return_value="Alice works at Example Labs. OCR invoice 7391.")
    rag.update_llm_role_config("vlm", model_func=vlm)
    image = Image.new("RGB", (500, 150), "white")
    ImageDraw.Draw(image).text((20, 40), "OCR invoice 7391", fill="black")
    stream = BytesIO()
    image.save(stream, format=format)
    response = await client.post(
        "/documents/upload",
        files={"file": (f"scan.{suffix}", stream.getvalue())},
        data={"folder_id": folder.id, "fast_index": "true"},
    )
    assert response.status_code == 200, response.text
    docs = await rag.doc_status.get_docs_by_track_id(response.json()["track_id"])
    doc_id, doc = next(iter(docs.items()))
    assert doc.status.value == "processed", doc.error_msg
    assert folder.id in doc.metadata["folder_ids"]
    assert "7391" in (await rag.full_docs.get_by_id(doc_id))["content"]
    assert doc.chunks_list
    vlm.assert_awaited_once()


async def test_vlm_failure_is_visible_and_later_upload_recovers(upload_app, monkeypatch):
    from PIL import Image
    from unittest.mock import AsyncMock

    client, rag, folder, _ = upload_app
    monkeypatch.setenv("LIGHTRAG_OCR_ENGINE", "vlm")
    rag.vlm_process_enable = True
    rag.update_llm_role_config("vlm", model_func=AsyncMock(side_effect=RuntimeError("OCR unavailable")))
    stream = BytesIO()
    Image.new("RGB", (40, 40)).save(stream, format="PNG")
    response = await client.post(
        "/documents/upload", files={"file": ("broken-ocr.png", stream.getvalue())},
        data={"folder_id": folder.id},
    )
    docs = await rag.doc_status.get_docs_by_track_id(response.json()["track_id"])
    doc = next(iter(docs.values()))
    assert doc.status.value == "failed"
    assert "OCR unavailable" in doc.error_msg
    assert folder.id in doc.metadata["folder_ids"]
    response = await client.post(
        "/documents/upload", files={"file": ("after-ocr.txt", b"A valid later document.")},
        data={"fast_index": "true"},
    )
    docs = await rag.doc_status.get_docs_by_track_id(response.json()["track_id"])
    assert next(iter(docs.values())).status.value == "processed"


@pytest.mark.parametrize("extension", ["doc", "xls"])
async def test_real_legacy_office_upload_and_repeat(upload_app, tmp_path, extension):
    """Exercise real legacy binary files, not renamed XML or mocked conversion."""
    import shutil
    import subprocess

    converter = shutil.which("libreoffice") or shutil.which("soffice")
    if not converter:
        pytest.skip("LibreOffice is required for real legacy Office conversion")
    if extension == "doc":
        from docx import Document
        original = tmp_path / "office-source.docx"
        document = Document()
        document.add_paragraph("Legacy office acceptance 6418")
        document.save(original)
    else:
        from openpyxl import Workbook
        original = tmp_path / "office-source.xlsx"
        workbook = Workbook()
        workbook.active.append(["Legacy office acceptance", "6418"])
        workbook.save(original)
    subprocess.run(
        [converter, f"-env:UserInstallation={(tmp_path / 'lo-profile').as_uri()}",
         "--headless", "--convert-to", extension, "--outdir", str(tmp_path), str(original)],
        check=True, capture_output=True, timeout=120,
    )
    content = (tmp_path / f"office-source.{extension}").read_bytes()
    assert content.startswith(bytes.fromhex("d0cf11e0a1b11ae1"))
    client, rag, folder, _ = upload_app
    ids = []
    for _ in range(2):
        response = await client.post(
            "/documents/upload",
            files={"file": (f"legacy.{extension.upper()}", content, "application/octet-stream")},
            data={"folder_id": folder.id, "fast_index": "true"},
        )
        assert response.status_code == 200, response.text
        docs = await rag.doc_status.get_docs_by_track_id(response.json()["track_id"])
        doc_id, doc = next(iter(docs.items()))
        assert doc.status.value == "processed", doc.error_msg
        assert "6418" in (await rag.full_docs.get_by_id(doc_id))["content"]
        assert folder.id in doc.metadata["folder_ids"]
        ids.append(doc_id)
    assert len(set(ids)) == 1
    docs = await rag.doc_status.get_docs_by_status(routes.DocStatus.PROCESSED)
    assert len(docs) == 1


async def test_new_upload_does_not_retry_previous_analysis_failure(upload_app, monkeypatch):
    from unittest.mock import AsyncMock

    client, rag, _, _ = upload_app
    await rag.apipeline_enqueue_documents(
        "Old content whose model analysis failed", file_paths="old-failed.txt",
    )
    old = await rag.doc_status.get_docs_by_status(routes.DocStatus.PENDING)
    old_id = next(iter(old))
    row = await rag.doc_status.get_by_id(old_id)
    row.update(status="failed", error_msg="Previous model analysis timed out")
    await rag.doc_status.upsert({old_id: row})
    model = AsyncMock(side_effect=AssertionError("Old failed document must not run"))
    rag.update_llm_role_config("extract", model_func=model)
    response = await client.post(
        "/documents/upload", files={"file": ("new-only.txt", b"New content should be indexed")},
        data={"fast_index": "true"},
    )
    assert response.status_code == 200
    docs = await rag.doc_status.get_docs_by_track_id(response.json()["track_id"])
    assert next(iter(docs.values())).status.value == "processed"
    model.assert_not_awaited()
    old_row = await rag.doc_status.get_by_id(old_id)
    assert old_row["status"] == "failed"
    assert old_row["error_msg"] == "Previous model analysis timed out"

    rag.update_llm_role_config("extract", model_func=fake_llm)
    response = await client.post("/documents/reprocess_failed")
    assert response.status_code == 200
    assert (await rag.doc_status.get_by_id(old_id))["status"] == "processed"


async def test_explicit_retry_requested_during_upload_is_not_lost(upload_app, monkeypatch):
    _, rag, _, _ = upload_app
    await rag.apipeline_enqueue_documents("Old failed analysis", file_paths="retry-old.txt")
    old_id = next(iter(await rag.doc_status.get_docs_by_status(routes.DocStatus.PENDING)))
    old = await rag.doc_status.get_by_id(old_id)
    old.update(status="failed", error_msg="Model error")
    await rag.doc_status.upsert({old_id: old})
    await rag.apipeline_enqueue_documents("New pending content", file_paths="retry-new.txt")
    original_batch = rag._run_pipeline_batch
    batches = []

    async def batch(documents, **kwargs):
        batches.append(set(documents))
        if len(batches) == 1:
            assert old_id not in documents
            # Model an explicit retry API call arriving while the upload is busy.
            await rag.apipeline_process_enqueue_documents()
        await original_batch(documents, **kwargs)

    monkeypatch.setattr(rag, "_run_pipeline_batch", batch)
    await rag.apipeline_process_enqueue_documents(retry_failed=False)
    assert len(batches) == 2
    assert old_id in batches[1]
    assert (await rag.doc_status.get_by_id(old_id))["status"] == "processed"
